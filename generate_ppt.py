from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Color Palette
BG_DARK     = RGBColor(0x08, 0x0D, 0x18)
ACCENT_CYAN = RGBColor(0x00, 0xD4, 0xFF)
ACCENT_PURP = RGBColor(0x8B, 0x5C, 0xF6)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_WHITE  = RGBColor(0xC8, 0xD8, 0xF0)
GOLD        = RGBColor(0xFF, 0xC8, 0x4A)
GREEN       = RGBColor(0x00, 0xE5, 0x8A)
RED_SOFT    = RGBColor(0xFF, 0x6B, 0x6B)
CARD_BG     = RGBColor(0x10, 0x18, 0x2C)
HEADER_BG   = RGBColor(0x0C, 0x12, 0x22)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

# ── Helpers ──────────────────────────────────────────────────────────────────
def bg(slide, color=BG_DARK):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color

def rect(slide, l, t, w, h, color, border=None, border_w=1):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if border:
        s.line.color.rgb = border
        s.line.width = Pt(border_w)
    else:
        s.line.fill.background()
    return s

def txt(slide, text, l, t, w, h, size=13, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, italic=False):
    b = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    b.word_wrap = True
    tf = b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.italic = italic
    return b

def line(slide, l, t, w, color=ACCENT_CYAN, h=0.045):
    rect(slide, l, t, w, h, color)

def card(slide, l, t, w, h, title, bullets,
         tc=ACCENT_CYAN, bc=SOFT_WHITE, bg_c=CARD_BG, border=ACCENT_PURP):
    rect(slide, l, t, w, h, bg_c, border, 1)
    txt(slide, title, l+0.15, t+0.1, w-0.25, 0.38, size=12, bold=True, color=tc)
    b = slide.shapes.add_textbox(Inches(l+0.15), Inches(t+0.5), Inches(w-0.25), Inches(h-0.6))
    b.word_wrap = True
    tf = b.text_frame; tf.word_wrap = True
    first = True
    for blt in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(2)
        r = p.add_run()
        r.text = f"  {blt}"
        r.font.size = Pt(10)
        r.font.color.rgb = bc

def header(slide, title, subtitle=None):
    rect(slide, 0, 0, 13.33, 0.88, HEADER_BG)
    line(slide, 0, 0.88, 13.33, ACCENT_CYAN, 0.05)
    txt(slide, title, 0.4, 0.13, 12.5, 0.55, size=20, bold=True, color=ACCENT_CYAN)
    if subtitle:
        txt(slide, subtitle, 0.4, 0.5, 12.5, 0.32, size=11, color=SOFT_WHITE)

def footer(slide, num, total=7):
    txt(slide, f"Slide {num} / {total}", 11.6, 7.12, 1.6, 0.28,
        size=9, color=SOFT_WHITE, align=PP_ALIGN.RIGHT)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title + Why We Need This
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl)

# left side bar
rect(sl, 0, 0, 0.55, 7.5, ACCENT_PURP)
rect(sl, 0.55, 0, 0.05, 7.5, ACCENT_CYAN)

# label
txt(sl, "VIVA PRESENTATION  |  2026", 0.85, 0.28, 12, 0.38,
    size=10, color=SOFT_WHITE)

# Brand name — SiddhiAI
# Siddhi = perfect mastery / accomplishment (Sanskrit)
txt(sl, "SiddhiAI", 0.85, 0.7, 12, 1.4,
    size=76, bold=True, color=ACCENT_CYAN)
txt(sl, "Siddhi (Sanskrit) — Perfect Mastery & Accomplishment", 0.85, 1.96, 12, 0.38,
    size=11, color=SOFT_WHITE, italic=True)
txt(sl, "AI-Powered Adaptive Interview Preparation Platform", 0.85, 2.38, 12, 0.42,
    size=18, bold=False, color=WHITE)

line(sl, 0.85, 2.9, 11.5, ACCENT_CYAN, 0.05)

# WHY WE NEED THIS
txt(sl, "Why Does This Platform Need to Exist?", 0.85, 3.05, 12, 0.4,
    size=14, bold=True, color=GOLD)

problems = [
    ("Talent Wasted at the Gate",
     "Thousands of highly capable engineers fail interviews not due to lack of knowledge, but lack of structured, personalised practice."),
    ("Generic Tools, Generic Failure",
     "Existing platforms use the same question banks for everyone — blind to your resume, your target company, your actual skill gaps."),
    ("No Real Feedback Loop",
     "Candidates receive a score but never understand WHY they failed or WHAT to fix. Growth becomes guesswork."),
    ("The Confidence Deficit",
     "Without realistic, adaptive mock interviews, even strong candidates freeze under actual interview pressure."),
]
y = 3.55
for bold_part, soft_part in problems:
    txt(sl, f"  {bold_part}  —  {soft_part}", 0.9, y, 11.9, 0.42,
        size=11, color=SOFT_WHITE)
    y += 0.44

line(sl, 0.85, 5.4, 11.5, ACCENT_PURP, 0.04)
txt(sl, "SiddhiAI bridges this gap — intelligent, resume-aware, company-calibrated, and brutally honest about where you stand.",
    0.85, 5.5, 12, 0.42, size=12, bold=True, color=GREEN)

footer(sl, 1)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Platform Overview
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl)
header(sl, "PLATFORM OVERVIEW",
       "What SiddhiAI is, who it is for, and what makes it different")

txt(sl, "What Is SiddhiAI?", 0.5, 1.08, 12, 0.38, size=15, bold=True, color=WHITE)
desc = (
    "SiddhiAI is an end-to-end, AI-powered interview preparation platform that reads your resume, "
    "identifies your real skill gaps, and generates personalised interview questions through Google Gemini. "
    "It evaluates your answers fairly — the way a real interviewer would — and builds you a custom "
    "learning roadmap to ensure you keep improving with every session."
)
txt(sl, desc, 0.5, 1.52, 12.3, 0.88, size=12, color=SOFT_WHITE)

# 3 stat blocks
for val, lbl, lx in [("Google Gemini", "Powered By", 0.5),
                      ("10 Modules", "Core Feature Set", 4.6),
                      ("Fair & Explainable", "Scoring Engine", 8.7)]:
    rect(sl, lx, 2.58, 3.8, 1.12, CARD_BG, ACCENT_PURP)
    txt(sl, val, lx+0.12, 2.65, 3.55, 0.52, size=22, bold=True,
        color=ACCENT_CYAN, align=PP_ALIGN.CENTER)
    txt(sl, lbl, lx+0.12, 3.15, 3.55, 0.38, size=11,
        color=SOFT_WHITE, align=PP_ALIGN.CENTER)

txt(sl, "Technology Stack", 0.5, 3.9, 12, 0.36, size=13, bold=True, color=ACCENT_PURP)
stack = [
    ("React + Vite", "Frontend"),
    ("FastAPI / Python", "Backend"),
    ("Google Gemini", "AI Engine"),
    ("SQLite / PostgreSQL", "Database"),
    ("Framer Motion", "Animations"),
]
sx = 0.5
for tech, role in stack:
    rect(sl, sx, 4.35, 2.3, 0.82, CARD_BG, ACCENT_CYAN, 1)
    txt(sl, tech, sx+0.08, 4.4, 2.12, 0.38, size=11, bold=True,
        color=ACCENT_CYAN, align=PP_ALIGN.CENTER)
    txt(sl, role, sx+0.08, 4.76, 2.12, 0.32, size=9.5,
        color=SOFT_WHITE, align=PP_ALIGN.CENTER)
    sx += 2.55

footer(sl, 2)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Core Modules (1–5)
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl)
header(sl, "CORE MODULES  (1 of 2) — Intelligence & Gap Analysis",
       "Modules 1 through 5: Building your complete candidate profile")

module_data_1 = [
    ("M1  Resume Intelligence",
     ["Parses PDF/DOCX/TXT resume via Gemini API",
      "Extracts skills, projects, education, experience",
      "Alias normalisation: ReactJS -> React, Postgres -> PostgreSQL",
      "Stored profile used across all downstream modules"]),
    ("M2  Job Description Intelligence",
     ["50+ companies pre-seeded with real interview data",
      "Stores: focus areas, common topics, difficulty tier",
      "Company hiring bar injected into every AI prompt",
      "Filter by company name, tier, or domain"]),
    ("M3  Skill Matching & Gap Engine",
     ["Compares resume skills vs. role-required skills",
      "Categories: Strong / Developing / Missing",
      "Computes Readiness Score (0-100%) per role",
      "Feeds critical gaps into question generation"]),
    ("M4  Skill Knowledge Graph",
     ["Visual map of skill relationships & prerequisites",
      "e.g. React -> JavaScript -> Node.js chain",
      "Identifies which skills unblock the most paths",
      "Foundation for the personalised roadmap"]),
    ("M5  Personalised Preparation Roadmap",
     ["Auto-generated from skill gaps after each session",
      "Priority levels: High / Medium / Low",
      "Each recommendation includes clear reason/context",
      "Evolves dynamically as candidate improves"]),
]

lx = 0.35
for i, (title, bullets) in enumerate(module_data_1):
    card(sl, lx, 1.08, 2.46, 5.55, title, bullets,
         tc=GOLD if i == 0 else ACCENT_CYAN)
    lx += 2.56

footer(sl, 3)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Core Modules (6–10)
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl)
header(sl, "CORE MODULES  (2 of 2) — AI Interview & Evaluation Engine",
       "Modules 6 through 10: The live interview, scoring, and growth loop")

module_data_2 = [
    ("M6  AI Question Generation",
     ["Gemini generates questions — NOT a static bank",
      "Anchored to your real resume projects & tools",
      "Company-calibrated: Google != Startup difficulty",
      "Two personas: Technical (Adinath) & Behavioural (Veda)"]),
    ("M7  Adaptive Mock Interview",
     ["5-round live interview session",
      "Multi-turn memory: each Q builds on previous answers",
      "Difficulty scales in real-time to your performance",
      "Auto-completes session & triggers scoring"]),
    ("M8  Multi-dim. Evaluation",
     ["4 dimensions: Score, Correctness, Relevance, Completeness",
      "Semantic scoring — not just keyword matching",
      "Gemini reads meaning, just like a real interviewer",
      "Detected vs. Missing concepts shown clearly"]),
    ("M9  Weakness Detection & Recs",
     ["Tracks missing concepts across all 5 rounds",
      "Generates 3 prioritised recommendations via Gemini",
      "Each rec: Topic + Priority + Reason",
      "Persistent across all future sessions"]),
    ("M10  Performance Dashboard",
     ["Interview history: scores, dates, roles",
      "Per-session: Q -> Answer -> Score -> Feedback",
      "Readiness % updates after every session",
      "Long-term progress & skill trend tracking"]),
]

lx = 0.35
for i, (title, bullets) in enumerate(module_data_2):
    card(sl, lx, 1.08, 2.46, 5.55, title, bullets,
         tc=GREEN if i == 2 else ACCENT_CYAN)
    lx += 2.56

footer(sl, 4)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — How It Works (Workflow)
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl)
header(sl, "HOW IT WORKS  —  End-to-End Workflow",
       "From resume upload to final performance report in 6 clear steps")

steps = [
    ("01", "Upload Resume",
     "Candidate uploads their CV. Gemini API reads and extracts skills, projects, and experience into a structured profile."),
    ("02", "Select Role & Company",
     "Pick your target job role and the company you are interviewing for. The system calibrates its entire engine to that hiring bar."),
    ("03", "Skill Gap Analysis",
     "SiddhiAI compares your resume against the role's required skills. You see exactly where you are strong and where you are missing."),
    ("04", "AI Mock Interview Begins",
     "5 adaptive rounds start. Each question is generated live by Gemini — tailored to your resume, your gaps, and your target company."),
    ("05", "Answer Evaluation",
     "Every answer is scored across 4 dimensions. You get a fair, semantic score with explicit feedback on what you said and what you missed."),
    ("06", "Roadmap & Report",
     "Session ends. Your overall score is calculated, weak topics are detected, and a personalised learning roadmap is ready immediately."),
]

sy = 1.08
for i, (num, title, desc) in enumerate(steps):
    col_x = 0.35 if i < 3 else 6.85
    row_y = sy + (i % 3) * 1.82

    rect(sl, col_x, row_y, 6.2, 1.65, CARD_BG, ACCENT_PURP, 1)
    rect(sl, col_x, row_y, 0.58, 1.65, ACCENT_PURP)
    txt(sl, num, col_x+0.02, row_y+0.52, 0.55, 0.55,
        size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(sl, title, col_x+0.68, row_y+0.1, 5.35, 0.42,
        size=12, bold=True, color=GOLD)
    txt(sl, desc, col_x+0.68, row_y+0.52, 5.35, 1.0,
        size=10.5, color=SOFT_WHITE)

footer(sl, 5)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Use Cases & Impact
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl)
header(sl, "USE CASES  &  REAL-WORLD IMPACT",
       "Who benefits from SiddhiAI and how it transforms career outcomes")

use_cases = [
    ("Fresh Graduates", ACCENT_CYAN,
     "First-time job seekers who need structured, company-specific practice before stepping into FAANG or top-tier campus interviews."),
    ("Career Switchers", ACCENT_PURP,
     "Professionals moving from one domain to another who need to quickly identify and close skill gaps for a new target role."),
    ("Working Professionals", GREEN,
     "Experienced engineers preparing for senior or staff-level interviews at high-growth companies without time to waste on generic prep."),
    ("Bootcamp / College Students", GOLD,
     "Students in CS or engineering programs who want to benchmark themselves against real industry standards before placements."),
]

bx = 0.38
for who, col, why in use_cases:
    rect(sl, bx, 1.08, 2.95, 2.45, CARD_BG, col, 1)
    rect(sl, bx, 1.08, 2.95, 0.42, col)
    txt(sl, who, bx+0.1, 1.1, 2.73, 0.36,
        size=12, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER)
    txt(sl, why, bx+0.12, 1.58, 2.7, 1.8,
        size=10.5, color=SOFT_WHITE)
    bx += 3.12

# Impact stats
rect(sl, 0.38, 3.72, 12.5, 0.05, ACCENT_CYAN)
txt(sl, "Measurable Impact on Candidate Performance", 0.38, 3.85, 12.5, 0.4,
    size=14, bold=True, color=WHITE)

impacts = [
    ("Resume-Grounded", "Questions anchored to YOUR real experience — not generic scenarios"),
    ("Company-Calibrated", "Difficulty tuned to actual hiring bars of target companies"),
    ("Fair Scoring", "Semantic evaluation that rewards understanding, not memorised phrases"),
    ("Continuous Growth", "Roadmap evolves after every session — improvement never stops"),
]

ix = 0.38
for stat, desc in impacts:
    rect(sl, ix, 4.38, 2.95, 1.42, CARD_BG, ACCENT_PURP, 1)
    txt(sl, stat, ix+0.12, 4.44, 2.7, 0.42,
        size=11.5, bold=True, color=ACCENT_CYAN, align=PP_ALIGN.CENTER)
    txt(sl, desc, ix+0.12, 4.9, 2.7, 0.82,
        size=9.5, color=SOFT_WHITE, align=PP_ALIGN.CENTER)
    ix += 3.12

footer(sl, 6)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Summary & Future Scope
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(prs.slide_layouts[6])
bg(sl)
header(sl, "SUMMARY  &  FUTURE SCOPE",
       "What we have built and where SiddhiAI goes next")

# Left: module summary
txt(sl, "All 10 Modules — Delivered", 0.4, 1.05, 6.2, 0.38,
    size=13, bold=True, color=WHITE)

mods = [
    "M1  Resume Intelligence",
    "M2  Job Description Intelligence",
    "M3  Skill Matching & Gap Engine",
    "M4  Skill Knowledge Graph",
    "M5  Personalised Preparation Roadmap",
    "M6  AI Question Generation (Gemini)",
    "M7  Adaptive Mock Interview (5 Rounds)",
    "M8  Multi-Dimensional Answer Evaluation",
    "M9  Weakness Detection + Recommendations",
    "M10 Performance Dashboard & Progress Tracking",
]
my = 1.52
for i, m in enumerate(mods):
    col = ACCENT_CYAN if i < 5 else GREEN
    txt(sl, f"  {m}", 0.5, my, 5.9, 0.36, size=10.5, color=col, bold=False)
    my += 0.38

# Right: future scope
rect(sl, 6.9, 1.05, 6.0, 5.55, CARD_BG, ACCENT_PURP, 1)
txt(sl, "Future Scope", 7.05, 1.12, 5.7, 0.38,
    size=13, bold=True, color=GOLD)

future = [
    ("Voice Interview Mode",
     "Real-time speech recognition + text-to-speech for a fully simulated voice interview round"),
    ("Live Coding Environment",
     "In-browser DSA coding sandbox with AI-generated algorithmic challenges per role"),
    ("JD Web Scraper",
     "Auto-fetch and analyse real job descriptions from LinkedIn, Naukri, or company portals"),
    ("Peer Mock Matching",
     "Two candidates interview each other — evaluated and scored by the AI as an observer"),
    ("Mobile App",
     "React Native app for on-the-go practice sessions, notifications, and daily prep streaks"),
]

fy = 1.58
for ftitle, fdesc in future:
    rect(sl, 6.95, fy, 5.85, 0.88, RGBColor(0x14, 0x1E, 0x35))
    txt(sl, ftitle, 7.08, fy+0.06, 5.6, 0.32, size=11, bold=True, color=ACCENT_CYAN)
    txt(sl, fdesc, 7.08, fy+0.38, 5.6, 0.44, size=9.5, color=SOFT_WHITE)
    fy += 0.98

# Final tagline
rect(sl, 0, 6.88, 13.33, 0.62, HEADER_BG)
txt(sl,
    "SiddhiAI  —  Siddhi (Sanskrit): Perfect Mastery.  Built to turn practice into accomplishment.",
    0.4, 6.93, 12.5, 0.38, size=12, bold=True,
    color=ACCENT_CYAN, align=PP_ALIGN.CENTER)

footer(sl, 7)

# ── Save ──────────────────────────────────────────────────────────────────────
out = r"c:\Users\kshit\OneDrive\Desktop\claude folder\SiddhiAI_Viva_PPT.pptx"
prs.save(out)
print(f"[DONE] Saved -> {out}")
